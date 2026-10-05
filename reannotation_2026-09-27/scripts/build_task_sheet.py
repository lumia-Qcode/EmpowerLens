"""
Build the week's task tracker (week of 28 Sep 2026).

    venv\\Scripts\\python.exe reannotation_2026-09-27/scripts/build_task_sheet.py

Writes ``planning/tasks_week_2026-09-28.xlsx``. Re-running overwrites the file,
so edit the TASKS list here rather than the spreadsheet if you want the script to
stay the source of truth; otherwise just edit the spreadsheet and leave this be.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "planning" / "tasks_week_2026-09-28.xlsx"

OWNERS = ["Nayab", "Lumia", "Izza", "Laiba", "Reesha", "Hurema", "Team", "Unassigned"]
STATUSES = ["Not started", "In progress", "Blocked", "Done"]

HEADERS = ["ID", "Module", "Task", "Detail", "Owner", "Deadline", "Status",
           "Depends on", "Notes"]

# Which part of the project each task belongs to.
MODULES = {
    "T01": "Re-annotation", "T02": "Re-annotation", "T03": "Re-annotation",
    "T04": "Re-annotation", "T05": "Re-annotation",
    "T06": "Data augmentation", "T07": "Data augmentation", "T08": "Data augmentation",
    "T09": "Transfer evaluation", "T10": "Transfer evaluation",
    "T11": "Span selection", "T12": "Span selection", "T13": "Span selection",
    "T14": "Model tuning", "T15": "Model tuning",
    "T16": "Entrepreneurial data", "T17": "Entrepreneurial data",
    "T18": "Entrepreneurial data", "T19": "Entrepreneurial data",
    "T20": "Entrepreneurial data", "T21": "Entrepreneurial data",
    "T22": "Specificity set", "T23": "Specificity set", "T24": "Specificity set",
}

MODULE_COLOURS = {
    "Re-annotation": "DDEBF7",
    "Data augmentation": "E2EFDA",
    "Transfer evaluation": "FFF2CC",
    "Span selection": "FCE4D6",
    "Model tuning": "E4DFEC",
    "Entrepreneurial data": "DAEEF3",
    "Specificity set": "F2DCDB",
}

# (id, task, detail, owner, deadline, depends_on, notes)
# Owners and the T09/T10 wording come from the team's reassignment (28 Sep).
# Rows are emitted sorted by deadline; IDs stay stable so dependencies still read.
TASKS = [
    ("T16", "Decide target dataset sizes",
     "How many seed + test items the entrepreneurial set needs",
     "Lumia", "Tue 29 Sep", "",
     "~100 non-distorted for +/-10% specificity CI; ~20/class -> 200-300 for per-class"),
    ("T09", "CBT-Bench: acquire + map labels + run on org dataset finetuned model",
     "Psychotherapy-LLM/CBT-Bench on HuggingFace; CBT-CD split, 146 items, 10 classes",
     "Izza", "Wed 30 Sep", "",
     "Write an explicit label mapping; record how many items had to be dropped"),
    ("T11", "Span approach: literature + decision",
     "BIO token tagging vs sentence classification vs extractive QA",
     "Lumia", "Wed 30 Sep", "",
     "Recommendation: sentence-level — annotators flagged sentences, report promises it"),
    ("T17", "Scraping: sources + ToS check",
     "Pick sources, confirm terms of service and redistribution limits",
     "Izza", "Wed 30 Sep", "T16",
     "Scraped text can be republished only as derived data in most cases"),
    ("T19", "Circulate form + start thread",
     "Post to relevant entrepreneurial groups/sites; collect responses",
     "Lumia", "Wed 30 Sep", "",
     "Consent + anonymisation first. Responses are EVALUATION ONLY, never prompts"),
    ("T12", "Build sentence labels from spans",
     "Split into sentences; label the one containing the gold span; use find_span offsets",
     "Lumia", "Fri 02 Oct", "T11",
     "Current span_crop/span_marked modes are ORACLE; the selector does not exist yet"),
    ("T14", "Hyperparameter tuning",
     "LR (2e-5/3e-5/5e-5), epochs, max_length, weight decay on the original dataset",
     "Izza", "Fri 02 Oct", "",
     "Val only. Confirm the winner across all 3 seeds; single-seed wins are noise"),
    ("T15", "Over/underfitting diagnosis",
     "Train vs val loss curves per epoch; early stop on val macro-F1",
     "Izza", "Fri 02 Oct", "T14",
     "~2,000 train rows: expect overfitting from epoch 3-4"),
    ("T01", "Label CS workbook",
     "annotator_4_CS_majors.xlsx — 413 rows (60 overlap + 293 training + 60 test)",
     "Team", "Sat 03 Oct", "",
     "Overlap 60 must be labelled independently, no discussion, or kappa is void"),
    ("T02", "Label psych workbooks",
     "annotator_1 Laiba 421 / annotator_2 Reesha 418 / annotator_3 Hurema 415 rows",
     "Team", "Sat 03 Oct", "",
     "~5h 10m each at 45 s/row; unfinished rows are acceptable"),
    ("T18", "Scrape entrepreneurial text",
     "Volume set by T16; cleaning and de-duplication included",
     "Izza", "Sat 03 Oct", "T17",
     "Unlabelled at this stage; labels come later with Laiba"),
    ("T20", "Triage form responses",
     "Keep responses usable for the entrepreneurial set",
     "Lumia", "Sat 03 Oct", "T19",
     "Held out at natural class distribution; do not balance"),
    ("T21", "Compile entrepreneurial dataset (unlabelled)",
     "Merge scraped + form text into one schema, ready for labelling",
     "Team", "Sat 03 Oct", "T18, T20",
     "WEEK GOAL. Labels follow later with Laiba and the team"),
    ("T03", "Merge labels + agreement",
     "Run scripts/relabel_agreement.py on returned workbooks",
     "Lumia", "Sun 04 Oct", "T01, T02",
     "Produces merged_labels.csv, agreement_report.md, unsure_for_adjudication.csv"),
    ("T22", "Seed the justified-concern set",
     "Harvest calibrated-pessimism rows from the relabelling output",
     "Nayab", "Mon 05 Oct", "T03",
     "Real, human-verified justified concerns; clinical topic, entrepreneurial rewrite needed"),
    ("T04", "Adjudicate Unsure rows",
     "Psych annotators resolve rows flagged Unsure",
     "Izza", "Tue 06 Oct", "T03",
     "Needed before the labels are final"),
    ("T05", "Separate errors from redefinitions",
     "Split label changes into genuine errors vs calibrated-pessimism redefinition",
     "Lumia", "Tue 06 Oct", "T03",
     "Original annotation had no justified-concern rule; keep the two countable apart"),
    ("T07", "Decide span handling for augmented rows",
     "Back-translation breaks 'Distorted part' offsets",
     "Lumia", "Wed 07 Oct", "T11",
     "MOVED BEFORE T06: it is a design decision augmentation depends on, not a follow-up"),
    ("T06", "Back-translation augmentation",
     "Inflate under-represented classes in the relabelled train split",
     "Lumia", "Thu 08 Oct", "T03, T04, T07",
     "TRAIN ONLY. Augmented rows stay in their source row's split or it is leakage"),
    ("T13", "Train + evaluate span selector",
     "Sentence accuracy and span overlap, reported separately from classification",
     "Lumia", "Thu 08 Oct", "T12",
     "Report the 0.402 oracle as an upper bound and the selector's cost as the gap"),
    ("T08", "Hand-check augmented samples",
     "30 rows per augmented class: did the distortion cue survive?",
     "Izza", "Fri 09 Oct", "T06",
     "'I'll never get funding' -> 'I might not' is no longer fortune-telling"),
    ("T23", "Write challenge-set pairs",
     "20 matched justified/distorted pairs per distortion family",
     "Team", "Fri 09 Oct", "T22",
     "Matched pairs give controls for free and block style shortcuts"),
    ("T10", "Evaluate fine-tuned model with relabelled+augmented dataset on CBT-Bench",
     "Transfer test of the relabelled + augmented model",
     "Izza", "Sat 10 Oct", "T06, T08, T09",
     "DATE MOVED from Fri 02 Oct: it needs the augmented data, which lands Thu 08 Oct"),
    ("T24", "Style-leak check on challenge set",
     "TF-IDF + logistic regression must score near chance on justified vs distorted",
     "Nayab", "Sat 10 Oct", "T23",
     "If it separates them easily (>0.75 F1) the set leaks surface cues — rewrite"),
]

FILLS = {
    "Not started": "FFFFFF",
    "In progress": "FFF2CC",
    "Blocked": "FFC7CE",
    "Done": "E2EFDA",
}


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"

    header_fill = PatternFill("solid", fgColor="1F4E5F")
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

    for tid, task, detail, owner, deadline, depends, note in TASKS:
        module = MODULES.get(tid, "")
        ws.append([tid, module, task, detail, owner, deadline, "Not started",
                   depends, note])
        cell = ws.cell(row=ws.max_row, column=2)
        cell.fill = PatternFill("solid", fgColor=MODULE_COLOURS.get(module, "FFFFFF"))
        cell.font = Font(bold=True)

    last = ws.max_row
    widths = [6, 20, 34, 62, 12, 13, 13, 14, 62]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for r in range(2, last + 1):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(
                vertical="top", wrap_text=c in (3, 4, 9)
            )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"

    dv_owner = DataValidation(type="list", formula1=f'"{",".join(OWNERS)}"',
                              allow_blank=True, showDropDown=False)
    dv_status = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"',
                               allow_blank=False, showDropDown=False)
    ws.add_data_validation(dv_owner)
    ws.add_data_validation(dv_status)
    dv_owner.add(f"E2:E{last}")
    dv_status.add(f"G2:G{last}")

    rng = f"A2:{get_column_letter(len(HEADERS))}{last}"
    for status, colour in FILLS.items():
        if status == "Not started":
            continue
        ws.conditional_formatting.add(
            rng,
            FormulaRule(formula=[f'$G2="{status}"'],
                        fill=PatternFill("solid", fgColor=colour), stopIfTrue=False),
        )

    # Legend / rules on a second tab.
    ws2 = wb.create_sheet("Notes")
    for row in [
        ["Week of Mon 28 Sep 2026 — deadline for the compiled dataset: Sat 03 Oct"],
        [""],
        ["Status values", ", ".join(STATUSES)],
        ["Modules", "Re-annotation, Data augmentation, Transfer evaluation, "
                    "Span selection, Model tuning, Entrepreneurial data, Specificity set"],
        ["Row colour", "amber = in progress, red = blocked, green = done"],
        [""],
        ["Standing rules"],
        ["1", "Augmentation is TRAIN ONLY; never val or test."],
        ["2", "Augmented rows stay in their source row's split, or it is leakage."],
        ["3", "Form/interview responses are EVALUATION ONLY — never prompts, never training."],
        ["4", "The test set is never balanced; it stays at its natural distribution."],
        ["5", "Thresholds are tuned on val. Test is read once, at the end."],
        ["6", "The challenge set is evaluation only and is frozen once verified."],
    ]:
        ws2.append(row)
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 80
    for r in ws2.iter_rows():
        for cell in r:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2["A1"].font = Font(bold=True)
    ws2["A6"].font = Font(bold=True)

    wb.save(OUT)
    print(f"wrote {OUT} ({last - 1} tasks)")


if __name__ == "__main__":
    main()
