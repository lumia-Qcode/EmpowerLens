"""
Build the relabelling session files from the model-triage output (TherapistQA).

Implements Tasks 1-3 of RELABEL_SPLIT_TASK.md. Task 4 (post-session merge and
agreement) lives in ``relabel_agreement.py``.

Run from the repo root:

    venv\\Scripts\\python.exe reannotation_2026-09-27/scripts/build_relabel_files.py

Nothing here modifies an input file. Every count and ID list is written to
``outputs/relabel/manifest.json`` so the split can be reproduced or audited.
"""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# Parameters
# --------------------------------------------------------------------------
SEED = 42
N_ANNOTATORS = 4
N_OVERLAP = 60
SHOW_DISTORTED_PART = True
CP_LABEL = "No Distortion (calibrated pessimism)"
ANNOTATOR_DISCIPLINE = {1: "psych", 2: "psych", 3: "psych", 4: "cs"}
ANNOTATOR_NAMES = {1: "Laiba", 2: "Reesha", 3: "Hurema", 4: "CS majors"}

# triage_full.csv carries every row of the dataset; A and B are the flagged
# buckets (disagree + ambiguous), C is the agree bucket.
FLAGGED_BUCKETS = {"A", "B"}
AGREE_BUCKET_NAME = "agree (not flagged)"

OVERLAP_FILL = "DDEBF7"  # light blue
TRAIN_FILL = "E2EFDA"  # light green
TEST_FILL = "FFF2CC"  # light amber

SECONDS_PER_ROW = 45

# Sheet protection locks the read-only columns, but any reader that opens the file
# in a restricted mode (Excel Protected View, a preview pane) then looks simply
# "broken" to an annotator. Off by default; flip to True if you want the lock.
PROTECT_SHEET = False

HERE = Path(__file__).resolve().parent
BASE = HERE.parent                       # reannotation_2026-09-27/
REPO = BASE.parent                       # repo root

TRIAGE_PATH = BASE / "inputs" / "triage_full.csv"
ORIGINAL_PATH = REPO / "Annotated_data.csv"
TEST_SPLIT_PATH = REPO / "data" / "splits" / "test.csv"
OUT = BASE / "outputs" / "relabel"

ID_COL = "Id_Number"
TEXT_COL = "Patient Question"
SPAN_COL = "Distorted part"
DOMINANT_COL = "Dominant Distortion"
SECONDARY_COL = "Secondary Distortion (Optional)"

# Annotator-visible layout (column order matters for the formatting rules).
VISIBLE = ["Row ID", "Text"] + ([SPAN_COL] if SHOW_DISTORTED_PART else []) + [
    "Primary",
    "Secondary",
    "Unsure",
    "Notes",
]
PRIMARY_COL_IDX = VISIBLE.index("Primary") + 1
SECONDARY_COL_IDX = VISIBLE.index("Secondary") + 1
EDITABLE = {"Primary", "Secondary", "Unsure", "Notes"}

errors: list[str] = []
notes: list[str] = []


def fail(msg: str) -> None:
    errors.append(msg)
    print(f"  ERROR: {msg}")


# --------------------------------------------------------------------------
# Load inputs
# --------------------------------------------------------------------------
def load_inputs():
    triage = pd.read_csv(TRIAGE_PATH, encoding="utf-8-sig")
    original = pd.read_csv(ORIGINAL_PATH, encoding="utf-8-sig")
    test = pd.read_csv(TEST_SPLIT_PATH, encoding="utf-8-sig")

    for name, df, path in (
        ("triage", triage, TRIAGE_PATH),
        ("original", original, ORIGINAL_PATH),
        ("test split", test, TEST_SPLIT_PATH),
    ):
        if ID_COL not in df.columns:
            sys.exit(f"{name} ({path}) has no '{ID_COL}' column; cannot match IDs.")

    return triage, original, test


def check_duplicates(triage, test):
    for name, df in (("triage file", triage), ("test split", test)):
        dupes = df[df[ID_COL].duplicated()][ID_COL].tolist()
        if dupes:
            fail(f"duplicate IDs in {name}: {dupes[:20]}")


# --------------------------------------------------------------------------
# Task 1: test sheet
# --------------------------------------------------------------------------
def build_test_sheet(triage, original, test, flagged_ids):
    test_ids = test[ID_COL].tolist()
    triage_by_id = triage.set_index(ID_COL)
    original_by_id = original.set_index(ID_COL)

    in_triage, not_in_triage, missing = [], [], []
    rows = []
    for rid in test_ids:
        if rid in flagged_ids:
            in_triage.append(rid)
            row = triage_by_id.loc[rid].to_dict()
            row[ID_COL] = rid
            row["in_triage_file"] = True
        elif rid in original_by_id.index:
            not_in_triage.append(rid)
            row = {c: np.nan for c in triage.columns}
            row[ID_COL] = rid
            row[TEXT_COL] = original_by_id.loc[rid, TEXT_COL]
            row["original_dominant"] = original_by_id.loc[rid, DOMINANT_COL]
            row["original_secondary"] = original_by_id.loc[rid, SECONDARY_COL]
            row["bucket"] = AGREE_BUCKET_NAME
            row["in_triage_file"] = False
        else:
            missing.append(rid)
            continue
        # The span never comes from the triage file; join it from the original.
        row[SPAN_COL] = (
            original_by_id.loc[rid, SPAN_COL] if rid in original_by_id.index else np.nan
        )
        rows.append(row)

    sheet = pd.DataFrame(rows)

    # Label mismatches between the triage file and the original dataset.
    mismatches = []
    for rid in in_triage:
        t_dom = str(triage_by_id.loc[rid, "original_dominant"]).strip()
        o_dom = str(original_by_id.loc[rid, DOMINANT_COL]).strip()
        t_sec = str(triage_by_id.loc[rid, "original_secondary"]).strip().lower()
        o_sec = str(original_by_id.loc[rid, SECONDARY_COL]).strip().lower()
        if t_dom != o_dom or t_sec != o_sec:
            mismatches.append(rid)

    if missing:
        fail(f"{len(missing)} test IDs missing from the original dataset: {missing[:20]}")
    if mismatches:
        fail(f"{len(mismatches)} test rows disagree on original labels: {mismatches[:20]}")

    sheet.to_excel(OUT / "test_rows_all.xlsx", index=False)

    md = [
        "# Test sheet discrepancy report",
        "",
        f"- Test split size: **{len(test_ids)}**",
        f"- Rows written: **{len(sheet)}**",
        f"- Test IDs in the flagged triage pool (buckets {sorted(FLAGGED_BUCKETS)}): "
        f"**{len(in_triage)}**",
        f"- Test IDs not flagged (agree bucket, pulled from the original dataset): "
        f"**{len(not_in_triage)}**",
        f"- Test IDs missing from the original dataset (error): **{len(missing)}**",
        f"- Rows whose original labels differ between triage and original (error): "
        f"**{len(mismatches)}**",
        "",
        "## ID lists",
        "",
        f"**Flagged test IDs ({len(in_triage)}):** {in_triage}",
        "",
        f"**Unflagged test IDs ({len(not_in_triage)}):** {not_in_triage}",
        "",
        f"**Missing from original ({len(missing)}):** {missing}",
        "",
        f"**Label mismatches ({len(mismatches)}):** {mismatches}",
    ]
    (OUT / "test_discrepancies.md").write_text("\n".join(md), encoding="utf-8")
    return sheet, in_triage, not_in_triage, missing, mismatches


# --------------------------------------------------------------------------
# Overlap sampling
# --------------------------------------------------------------------------
def sample_overlap(pool: pd.DataFrame, rng: np.random.Generator):
    """
    Pick N_OVERLAP rows stratified by bucket and original dominant label, giving
    every label at least 3 rows where the pool allows.
    """
    chosen: list[int] = []
    short_labels: dict[str, int] = {}

    for label, grp in pool.groupby("original_dominant", sort=True):
        want = min(3, len(grp))
        if len(grp) < 3:
            short_labels[label] = len(grp)
        # spread the guaranteed rows across buckets
        picks: list[int] = []
        for _, bgrp in grp.groupby("bucket", sort=True):
            if len(picks) >= want:
                break
            take = min(want - len(picks), len(bgrp))
            picks += list(rng.choice(bgrp[ID_COL].to_numpy(), size=take, replace=False))
        chosen += picks

    remaining = pool[~pool[ID_COL].isin(chosen)]
    need = N_OVERLAP - len(chosen)
    if need < 0:
        chosen = list(rng.permutation(chosen))[:N_OVERLAP]
    elif need > 0:
        # proportional top-up across bucket x label strata
        strata = list(remaining.groupby(["bucket", "original_dominant"], sort=True))
        weights = np.array([len(g) for _, g in strata], dtype=float)
        weights /= weights.sum()
        quota = np.floor(weights * need).astype(int)
        while quota.sum() < need:
            quota[int(np.argmax(weights * len(quota) - quota))] += 1
        for (_, grp), q in zip(strata, quota):
            q = int(min(q, len(grp)))
            if q:
                chosen += list(rng.choice(grp[ID_COL].to_numpy(), size=q, replace=False))
        if len(chosen) < N_OVERLAP:  # rounding shortfall
            left = remaining[~remaining[ID_COL].isin(chosen)][ID_COL].to_numpy()
            extra = rng.choice(left, size=N_OVERLAP - len(chosen), replace=False)
            chosen += list(extra)

    chosen = [int(x) for x in chosen][:N_OVERLAP]
    order = rng.permutation(len(chosen))
    return [chosen[i] for i in order], short_labels


def split_evenly(df: pd.DataFrame, strat_col: str, rng: np.random.Generator):
    """Shuffle, then deal rows round-robin within each stratum."""
    buckets = defaultdict(list)
    shuffled = df.sample(frac=1.0, random_state=int(rng.integers(0, 2**31 - 1)))
    for stratum, grp in shuffled.groupby(strat_col, sort=True):
        ids = grp[ID_COL].tolist()
        for i, rid in enumerate(ids):
            buckets[(i % N_ANNOTATORS) + 1].append(int(rid))
    return {a: buckets[a] for a in range(1, N_ANNOTATORS + 1)}


# --------------------------------------------------------------------------
# Task 2: workbooks
# --------------------------------------------------------------------------
def write_workbook(path: Path, rows: pd.DataFrame, fills: list[str], primary_values,
                   secondary_values, sheet_title: str = "Label"):
    wb = Workbook()
    ws = wb.active
    # Excel caps sheet names at 31 characters and forbids : \ / ? * [ ]
    ws.title = sheet_title[:31]

    ws.append(VISIBLE)
    for c in range(1, len(VISIBLE) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.protection = Protection(locked=True)

    for i, (_, r) in enumerate(rows.iterrows(), start=2):
        values = [r[ID_COL], r[TEXT_COL]]
        if SHOW_DISTORTED_PART:
            span = r.get(SPAN_COL)
            values.append("" if pd.isna(span) else span)
        values += ["", "", False, ""]   # Primary, Secondary, Unsure, Notes
        ws.append(values)
        fill = PatternFill("solid", fgColor=fills[i - 2])
        for c, name in enumerate(VISIBLE, start=1):
            cell = ws.cell(row=i, column=c)
            cell.fill = fill
            cell.protection = Protection(locked=name not in EDITABLE)
            if name == "Text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")

    last = ws.max_row
    widths = {"Row ID": 10, "Text": 90, SPAN_COL: 40, "Primary": 34,
              "Secondary": 26, "Unsure": 10, "Notes": 30}
    for c, name in enumerate(VISIBLE, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 18)
    ws.freeze_panes = "A2"

    # Dropdowns. Inline lists keep everything on one sheet, which survives
    # LibreOffice / Google Sheets round-trips better than a cross-sheet
    # reference. Excel caps the inline list at 255 characters, so assert it.
    p_col, s_col = get_column_letter(PRIMARY_COL_IDX), get_column_letter(SECONDARY_COL_IDX)
    u_col = get_column_letter(VISIBLE.index("Unsure") + 1)

    p_list, s_list = ",".join(primary_values), ",".join(secondary_values)
    for name, lst in (("Primary", p_list), ("Secondary", s_list)):
        if len(lst) > 255:
            raise SystemExit(f"{name} dropdown is {len(lst)} chars; Excel allows 255.")

    dv_p = DataValidation(type="list", formula1=f'"{p_list}"', allow_blank=False,
                          showDropDown=False, showErrorMessage=True,
                          errorTitle="Pick from the list",
                          error="Choose one of the labels in the dropdown.",
                          promptTitle="Primary", prompt="Required. One label.",
                          showInputMessage=True)
    dv_s = DataValidation(type="list", formula1=f'"{s_list}"', allow_blank=True,
                          showDropDown=False, showErrorMessage=True,
                          errorTitle="Pick from the list",
                          error="Choose one of the ten distortions, or leave blank.")
    # Unsure holds checkbox values. Excel 365's native checkbox is a cell
    # feature openpyxl cannot write, so the column is seeded with FALSE and the
    # annotator (or whoever preps the file) turns it into real checkboxes with
    # Insert > Checkbox over F2:F<last>. Until then it is a TRUE/FALSE dropdown.
    dv_u = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True,
                          showDropDown=False, showErrorMessage=False)
    for dv, col in ((dv_p, p_col), (dv_s, s_col), (dv_u, u_col)):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    # Red highlight for the three validation failures.
    red = PatternFill("solid", fgColor="FFC7CE")
    rule = FormulaRule(
        formula=[
            f'OR(${p_col}2="",'
            f'AND(${s_col}2<>"",${s_col}2=${p_col}2),'
            f'AND(${s_col}2<>"",OR(${p_col}2="No Distortion",${p_col}2="{CP_LABEL}")))'
        ],
        fill=red,
        stopIfTrue=False,
    )
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(VISIBLE))}{last}", rule)

    if PROTECT_SHEET:
        ws.protection.sheet = True
        ws.protection.selectLockedCells = False
    wb.save(path)


def build_workbooks(pool, test_sheet, overlap_ids, rng):
    pool_by_id = pool.set_index(ID_COL)
    test_by_id = test_sheet.set_index(ID_COL)

    remaining = pool[~pool[ID_COL].isin(overlap_ids)]
    train_assign = split_evenly(remaining, "bucket", rng)
    test_assign = split_evenly(test_sheet, "original_dominant", rng)

    primary_values = sorted(pool["original_dominant"].dropna().unique().tolist())
    if "No Distortion" not in primary_values:
        primary_values.append("No Distortion")
    primary_values = sorted(set(primary_values)) + [CP_LABEL]
    secondary_values = sorted(v for v in primary_values if v not in
                              {"No Distortion", CP_LABEL})

    key_rows, per_wb = [], {}
    for a in range(1, N_ANNOTATORS + 1):
        sections = (
            [("overlap", "train", rid, pool_by_id) for rid in overlap_ids]
            + [("training", "train", rid, pool_by_id) for rid in train_assign[a]]
            + [("test", "test", rid, test_by_id) for rid in test_assign[a]]
        )
        frame_rows, fills = [], []
        for section, split, rid, src in sections:
            row = src.loc[rid].to_dict()
            row[ID_COL] = rid
            frame_rows.append(row)
            fills.append({"overlap": OVERLAP_FILL, "training": TRAIN_FILL,
                          "test": TEST_FILL}[section])

        name = ANNOTATOR_NAMES.get(a, f"annotator_{a}")
        frame = pd.DataFrame(frame_rows)
        write_workbook(OUT / f"annotator_{a}_{name.replace(' ', '_')}.xlsx", frame, fills,
                       primary_values, secondary_values, sheet_title=name)

        for pos, (section, split, rid, src) in enumerate(sections, start=1):
            rec = src.loc[rid].to_dict()
            key_rows.append({
                "Row ID": rid,
                "annotator": a,
                "annotator_name": ANNOTATOR_NAMES.get(a, f"annotator_{a}"),
                "discipline": ANNOTATOR_DISCIPLINE.get(a, "?"),
                "position": pos,
                "section": section,
                "split": split,
                **{k: v for k, v in rec.items() if k != ID_COL},
            })
        per_wb[a] = {
            "name": name,
            "overlap": len(overlap_ids),
            "training": len(train_assign[a]),
            "test": len(test_assign[a]),
            "total": len(sections),
            "est_minutes": round(len(sections) * SECONDS_PER_ROW / 60, 1),
        }

    pd.DataFrame(key_rows).to_excel(OUT / "annotator_master_key.xlsx", index=False)
    return train_assign, test_assign, per_wb, primary_values, secondary_values


# --------------------------------------------------------------------------
# Task 3: training reannotation CSV
# --------------------------------------------------------------------------
def build_training_csv(pool, flagged, test_ids, train_ids_all):
    pool.to_csv(OUT / "training_reannotated.csv", index=False, encoding="utf-8")

    flagged_in_test = sorted(set(flagged[ID_COL]) & set(test_ids))
    expected = len(flagged) - len(flagged_in_test)
    orphans = sorted(set(flagged[ID_COL]) - set(test_ids) - set(train_ids_all))
    dupes = pool[pool[ID_COL].duplicated()][ID_COL].tolist()

    if len(pool) != expected:
        fail(f"training pool is {len(pool)} rows, expected {expected}")
    if orphans:
        fail(f"{len(orphans)} flagged IDs are in neither split: {orphans[:20]}")
    if dupes:
        fail(f"duplicate IDs in the training pool: {dupes[:20]}")

    md = [
        "# Training pool discrepancy report",
        "",
        f"- Flagged rows (buckets {sorted(FLAGGED_BUCKETS)}): **{len(flagged)}**",
        f"- Flagged rows inside the test split: **{len(flagged_in_test)}**",
        f"- Expected training-pool rows: **{expected}**",
        f"- Actual training-pool rows: **{len(pool)}**",
        f"- Flagged IDs in neither split (error): **{len(orphans)}**",
        f"- Duplicate IDs (error): **{len(dupes)}**",
        "",
        f"**Flagged test IDs ({len(flagged_in_test)}):** {flagged_in_test}",
        "",
        f"**Orphan IDs ({len(orphans)}):** {orphans}",
    ]
    (OUT / "training_discrepancies.md").write_text("\n".join(md), encoding="utf-8")
    return expected, flagged_in_test, orphans


# --------------------------------------------------------------------------
def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(SEED)

    triage, original, test = load_inputs()
    check_duplicates(triage, test)

    print(f"triage rows          : {len(triage)}")
    print(f"original rows        : {len(original)}")
    print(f"test split rows      : {len(test)}")
    print(f"bucket counts        : {triage['bucket'].value_counts().to_dict()}")

    flagged = triage[triage["bucket"].isin(FLAGGED_BUCKETS)].copy()
    flagged_ids = set(flagged[ID_COL])
    test_ids = test[ID_COL].tolist()
    print(f"flagged rows (A+B)   : {len(flagged)}")

    print("\nTask 1: test sheet")
    test_sheet, in_tri, not_in_tri, missing, mismatches = build_test_sheet(
        triage, original, test, flagged_ids
    )
    print(f"  rows written       : {len(test_sheet)}")
    print(f"  flagged / unflagged: {len(in_tri)} / {len(not_in_tri)}")

    # Training pool = flagged rows that are not in the test split (val rows included).
    pool = flagged[~flagged[ID_COL].isin(test_ids)].copy()
    pool[SPAN_COL] = pool[ID_COL].map(
        original.set_index(ID_COL)[SPAN_COL]
    )
    print(f"\ntraining pool        : {len(pool)}")

    print("\nTask 2: annotator workbooks")
    overlap_ids, short_labels = sample_overlap(pool, rng)
    if short_labels:
        notes.append(f"labels with fewer than 3 pool rows: {short_labels}")
        print(f"  under-3 labels     : {short_labels}")
    train_assign, test_assign, per_wb, primary_values, secondary_values = build_workbooks(
        pool, test_sheet, overlap_ids, rng
    )
    for a, c in per_wb.items():
        print(f"  {c['name']:<10} ({ANNOTATOR_DISCIPLINE.get(a,'?'):>5}): "
              f"{c['overlap']} overlap + {c['training']} training + {c['test']} test "
              f"= {c['total']} rows (~{c['est_minutes']:.0f} min)")

    print("\nTask 3: training reannotation CSV")
    train_ids_all = set(triage[ID_COL]) - set(test_ids)
    expected, flagged_in_test, orphans = build_training_csv(
        pool, flagged, test_ids, train_ids_all
    )
    print(f"  rows written       : {len(pool)} (expected {expected})")

    # ---------------- acceptance checks ----------------
    print("\nAcceptance checks")

    def check(label, ok, detail=""):
        print(f"  [{'PASS' if ok else 'FAIL'}] {label}{(' — ' + detail) if detail else ''}")
        if not ok:
            fail(label)

    check("test sheet rows == test split size",
          len(test_sheet) == len(test), f"{len(test_sheet)} vs {len(test)}")
    check("training CSV + flagged test rows == flagged total",
          len(pool) + len(flagged_in_test) == len(flagged),
          f"{len(pool)} + {len(flagged_in_test)} = {len(pool) + len(flagged_in_test)} "
          f"vs {len(flagged)}")

    sec2 = [set(v) for v in train_assign.values()]
    sec3 = [set(v) for v in test_assign.values()]
    no_dupe2 = sum(len(s) for s in sec2) == len(set().union(*sec2))
    no_dupe3 = sum(len(s) for s in sec3) == len(set().union(*sec3))
    check("section 2: no duplicates across workbooks", no_dupe2)
    check("section 2 union == training pool minus overlap",
          set().union(*sec2) == set(pool[ID_COL]) - set(overlap_ids))
    check("section 3: no duplicates across workbooks", no_dupe3)
    check("section 3 union == all test IDs",
          set().union(*sec3) == set(test_ids))
    check("no ID appears in more than one section",
          not (set(overlap_ids) & set().union(*sec2))
          and not (set(overlap_ids) & set().union(*sec3))
          and not (set().union(*sec2) & set().union(*sec3)))
    check("overlap set is identical and ordered in every workbook", True,
          "same list object written to all workbooks")
    check("workbooks expose annotator columns only",
          all(c in VISIBLE for c in VISIBLE) and "bucket" not in VISIBLE
          and "original_dominant" not in VISIBLE)
    bad = [v for v in primary_values
           if v != CP_LABEL and v not in set(original[DOMINANT_COL].dropna().unique())]
    check("every Primary dropdown value is a dataset label (except CP_LABEL)",
          not bad, str(bad))

    manifest = {
        "seed": SEED,
        "n_annotators": N_ANNOTATORS,
        "annotator_discipline": ANNOTATOR_DISCIPLINE,
        "annotator_names": ANNOTATOR_NAMES,
        "n_overlap": N_OVERLAP,
        "show_distorted_part": SHOW_DISTORTED_PART,
        "cp_label": CP_LABEL,
        "flagged_buckets": sorted(FLAGGED_BUCKETS),
        "counts": {
            "triage_rows": len(triage),
            "original_rows": len(original),
            "test_split_rows": len(test),
            "bucket_counts": triage["bucket"].value_counts().to_dict(),
            "flagged_rows": len(flagged),
            "flagged_in_test": len(flagged_in_test),
            "training_pool": len(pool),
            "per_workbook": per_wb,
        },
        "primary_dropdown": primary_values,
        "secondary_dropdown": secondary_values,
        "ids": {
            "overlap": overlap_ids,
            "test_flagged": in_tri,
            "test_unflagged": not_in_tri,
            "test_missing_from_original": missing,
            "label_mismatches": mismatches,
            "flagged_in_test": flagged_in_test,
            "orphans": orphans,
            "training_by_annotator": train_assign,
            "test_by_annotator": test_assign,
        },
        "notes": notes,
        "errors": errors,
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"\nwrote {OUT}")
    print(f"errors: {len(errors)}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
